"""
tests/test_docx_parser.py
Unit tests for the DOCX parser.
"""
import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import tempfile
import pytest
from pathlib import Path
from docx import Document as DocxDoc

from src.parsers.docx_parser import DocxParser
from src.parsers.document_model import DocumentModel


def _make_docx(paragraphs: list, headings: list = None) -> Path:
    """Helper: creates a minimal .docx in a temp file and returns its path."""
    doc = DocxDoc()
    if headings:
        for level, text in headings:
            doc.add_heading(text, level=level)
    for text in paragraphs:
        doc.add_paragraph(text)
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    doc.save(tmp.name)
    return Path(tmp.name)


class TestDocxParser:

    def test_returns_document_model(self):
        path = _make_docx(["Hello world"], headings=[(1, "Intro")])
        doc = DocxParser(path).parse()
        assert isinstance(doc, DocumentModel)

    def test_extracts_headings(self):
        path = _make_docx(
            ["Body text"],
            headings=[(1, "Title"), (2, "Sub-section")]
        )
        doc = DocxParser(path).parse()
        assert len(doc.headings) == 2
        assert doc.headings[0].level == 1
        assert doc.headings[0].text == "Title"

    def test_extracts_paragraphs(self):
        path = _make_docx(["Para one.", "Para two."])
        doc = DocxParser(path).parse()
        texts = [p.text for p in doc.paragraphs]
        assert "Para one." in texts
        assert "Para two." in texts

    def test_empty_document(self):
        path = _make_docx([])
        doc = DocxParser(path).parse()
        assert doc.word_count_from_raw == 0 or doc.raw_text == ""

    def test_page_count_positive(self):
        path = _make_docx(["word " * 400])
        doc = DocxParser(path).parse()
        assert doc.page_count >= 1

    def test_title_fallback_to_first_h1(self):
        path = _make_docx(["body"], headings=[(1, "My Document")])
        doc = DocxParser(path).parse()
        assert doc.title == "My Document"

    def test_file_type(self):
        path = _make_docx(["text"])
        doc = DocxParser(path).parse()
        assert doc.file_type == "docx"
